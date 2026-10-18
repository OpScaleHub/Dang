#!/usr/bin/env python3
"""Generates the Dang (دنگ) accounting prototype workbook (Persian, RTL).

Usage: python build_workbook.py [output.xlsx]

All sample data is SYNTHETIC (نمونه). The workbook is a validation prototype;
the production calculation engine is the Kotlin domain package. The same
scenario is asserted in app/src/test/.../WorkbookScenarioTest.kt.

The script also contains an independent exact-arithmetic (Fraction) reference
implementation of the allocation rules; its results are written to
expected-results.json and compared against the workbook's formula results by
verify_workbook.py.
"""
import json
import sys
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Tahoma"
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_CALC = Font(name=FONT, size=10, color="000000")
F_LINK = Font(name=FONT, size=10, color="008000")
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="37474F")
FILL_INPUT = PatternFill("solid", fgColor="FFF9C4")
FILL_CALC = PatternFill("solid", fgColor="FFFFFF")
FILL_CHECK = PatternFill("solid", fgColor="E8F5E9")
THIN = Side(style="thin", color="BDBDBD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0;[Red]-#,##0;"-"'
M3 = '#,##0.000'
INT = '#,##0'

YES, NO = "بله", "خیر"
B_EQUAL, B_OCC, B_AREA, B_CONS, B_CONS_COMMON = "مساوی", "نفر", "متراژ", "مصرف", "مصرف + سهم مشترک"
B_WEIGHT, B_EXPLICIT, B_PARK_ALL, B_PARK_USE, B_NONE = "وزنی", "مبلغ صریح", "پارکینگ (همه)", "پارکینگ (در استفاده)", "بدون تسهیم (صندوق)"
BASES = [B_EQUAL, B_OCC, B_AREA, B_CONS, B_CONS_COMMON, B_WEIGHT, B_EXPLICIT, B_PARK_ALL, B_PARK_USE, B_NONE]
FUND_OP, FUND_RES = "عملیاتی", "اندوخته"
FUNDS = [FUND_OP, FUND_RES]
RELATIONS = ["مالکیت", "تخصیص", "اجاره"]
ROLE_MAIN, ROLE_UNIT = "اصلی", "واحد"
ADJ_CREDIT, ADJ_DEBIT, ADJ_REFUND = "بستانکار (کاهش بدهی)", "بدهکار (افزایش بدهی)", "استرداد وجه"
ADJ_TYPES = [ADJ_CREDIT, ADJ_DEBIT, ADJ_REFUND]
METHODS = ["نقد", "کارت به کارت", "واریز بانکی", "چک"]

# --------------------------------------------------------------------------
# Synthetic sample data
# --------------------------------------------------------------------------
UNITS = [  # id, floor, number, area m², occupants (SAMPLE), sort
    ("1-1", 1, 1, 60, 2), ("1-2", 1, 2, 60, 1), ("1-3", 1, 3, 70, 3),
    ("2-1", 2, 1, 156, 4), ("2-2", 2, 2, 100, 2),
    ("3-1", 3, 1, 156, 5), ("3-2", 3, 2, 100, 3),
    ("4-1", 4, 1, 156, 4), ("4-2", 4, 2, 100, 2),
]
UNIT_IDS = [u[0] for u in UNITS]
PARKING = [  # space, unit, relation, in use, vehicle
    ("P1", "2-1", "مالکیت", YES, "خودرو نمونه A"), ("P2", "2-1", "مالکیت", NO, ""),
    ("P3", "3-1", "مالکیت", YES, "خودرو نمونه B"), ("P4", "3-1", "مالکیت", YES, "خودرو نمونه C"),
    ("P5", "4-1", "مالکیت", YES, "خودرو نمونه D"), ("P6", "2-2", "تخصیص", YES, "خودرو نمونه E"),
    ("P7", "3-2", "اجاره", YES, "خودرو نمونه F"), ("P8", "", "", NO, "بدون تخصیص"),
]
METERS = [  # id, role, unit, previous m³, current m³, rollover?, limit, note
    ("M-MAIN", ROLE_MAIN, "", 5120.0, 5262.0, NO, None, "کنتور اصلی ساختمان"),
    ("M-11", ROLE_UNIT, "1-1", 310.2, 322.7, NO, None, ""),
    ("M-12", ROLE_UNIT, "1-2", 205.0, 205.0, NO, None, "مصرف صفر (نمونه)"),
    ("M-13", ROLE_UNIT, "1-3", 9998.5, 3.2, YES, 10000, "گردش کنتور از ۹۹۹۹ به صفر"),
    ("M-21", ROLE_UNIT, "2-1", 1402.3, 1421.1, NO, None, ""),
    ("M-22-OLD", ROLE_UNIT, "2-2", 812.0, 815.4, NO, None, "کنتور قدیم - قرائت نهایی هنگام تعویض"),
    ("M-22-NEW", ROLE_UNIT, "2-2", 0.0, 6.1, NO, None, "کنتور جدید - قرائت اولیه صفر"),
    ("M-31", ROLE_UNIT, "3-1", 1650.0, 1672.4, NO, None, ""),
    ("M-32", ROLE_UNIT, "3-2", 980.5, 994.2, NO, None, ""),
    ("M-41", ROLE_UNIT, "4-1", 1710.0, 1731.3, NO, None, ""),
    ("M-42", ROLE_UNIT, "4-2", 1005.0, 1017.6, NO, None, ""),
]
CATEGORIES = [  # name, default basis, fund, water?
    ("نظافت", B_OCC, FUND_OP, NO), ("هزینه‌های عمومی", B_EQUAL, FUND_OP, NO),
    ("تعمیر و نگهداری ساختمان", B_AREA, FUND_OP, NO), ("آب", B_CONS_COMMON, FUND_OP, YES),
    ("برق پارکینگ", B_PARK_USE, FUND_OP, NO), ("سرویس آسانسور", B_WEIGHT, FUND_OP, NO),
    ("تعمیرات اختصاصی", B_EXPLICIT, FUND_OP, NO), ("تعمیرات اضطراری", B_NONE, FUND_RES, NO),
    ("اندوخته", B_EQUAL, FUND_RES, NO), ("برق مشاعات", B_AREA, FUND_OP, NO),
    ("لوله‌کشی", B_EQUAL, FUND_OP, NO), ("فضای سبز", B_EQUAL, FUND_OP, NO),
    ("امنیت", B_EQUAL, FUND_OP, NO), ("هزینه‌های اداری", B_EQUAL, FUND_OP, NO),
    ("سایر", B_EQUAL, FUND_OP, NO),
]
EXPENSES = [  # id, date, desc, category, supplier, amount, basis, fund
    ("E1", "1405/07/05", "نظافت راه‌پله - مهر", "نظافت", "شرکت خدماتی نمونه", 9_000_000, B_OCC, FUND_OP),
    ("E2", "1405/07/05", "هزینه‌های عمومی ماهانه", "هزینه‌های عمومی", "", 4_500_000, B_EQUAL, FUND_OP),
    ("E3", "1405/07/12", "سرویس موتورخانه و نگهداری", "تعمیر و نگهداری ساختمان", "تاسیسات نمونه", 12_000_000, B_AREA, FUND_OP),
    ("E4", "1405/07/28", "قبض آب دوره", "آب", "شرکت آب (نمونه)", 18_000_000, B_CONS_COMMON, FUND_OP),
    ("E5", "1405/07/28", "قبض برق پارکینگ", "برق پارکینگ", "شرکت برق (نمونه)", 1_600_000, B_PARK_USE, FUND_OP),
    ("E6", "1405/07/15", "سرویس ماهانه آسانسور", "سرویس آسانسور", "آسانسور نمونه", 6_000_000, B_WEIGHT, FUND_OP),
    ("E7", "1405/07/20", "تعمیر درب ورودی واحدها", "تعمیرات اختصاصی", "تعمیرکار نمونه", 2_500_000, B_EXPLICIT, FUND_OP),
    ("E8", "1405/07/22", "تعمیر اضطراری لوله اصلی", "تعمیرات اضطراری", "لوله‌کش نمونه", 3_000_000, B_NONE, FUND_RES),
]
EXP_CAP = 12
CHARGE_ITEMS = [  # id, name, category, fund, amount, basis
    ("C1", "حق اندوخته ماهانه", "اندوخته", FUND_RES, 9_000_000, B_EQUAL),
]
CHG_CAP = 5
# weights (elevator: floor-based SAMPLE weights) and explicit amounts
WEIGHTS = {
    "E6": {"1-1": 0, "1-2": 0, "1-3": 0, "2-1": 1, "2-2": 1, "3-1": 2, "3-2": 2, "4-1": 3, "4-2": 3},
    "E7": {"2-1": 1_500_000, "4-2": 1_000_000},
}
OPENING = {"cash": 60_000_000, "reserve": 40_000_000, "payable": 0,
           "units": {"1-2": 1_200_000, "3-1": -500_000}}
PAYMENTS = [  # date, unit, amount, method, ref
    ("1405/07/25", "1-1", 3_000_000, "کارت به کارت", "S-001"),
    ("1405/07/26", "1-2", 2_000_000, "نقد", "S-002"),
    ("1405/07/26", "2-1", 9_000_000, "واریز بانکی", "S-003"),
    ("1405/07/27", "2-2", 4_500_000, "کارت به کارت", "S-004"),
    ("1405/07/27", "3-1", 10_000_000, "واریز بانکی", "S-005"),
    ("1405/07/28", "3-2", 4_000_000, "نقد", "S-006"),
    ("1405/07/29", "4-1", 5_000_000, "کارت به کارت", "S-007"),
    ("1405/07/30", "4-2", 4_600_000, "واریز بانکی", "S-008"),
]
PAY_CAP = 30
SUPPLIER_PAYMENTS = [  # date, expense id, amount, ref
    ("1405/07/06", "E1", 9_000_000, "SP-01"), ("1405/07/06", "E2", 4_500_000, "SP-02"),
    ("1405/07/14", "E3", 6_000_000, "SP-03"), ("1405/07/29", "E5", 1_600_000, "SP-04"),
    ("1405/07/16", "E6", 6_000_000, "SP-05"), ("1405/07/21", "E7", 2_500_000, "SP-06"),
    ("1405/07/22", "E8", 3_000_000, "SP-07"),
]
SP_CAP = 30
ADJUSTMENTS = [  # date, unit, type, amount, fund, reason
    ("1405/07/30", "1-2", ADJ_CREDIT, 300_000, FUND_OP, "اصلاح نمونه: کسر هزینه نظافت دوره قبل"),
    ("1405/07/30", "3-1", ADJ_REFUND, 500_000, FUND_OP, "استرداد بستانکاری افتتاحیه (نمونه)"),
]
ADJ_CAP = 20
TRANSFERS = [("1405/07/30", FUND_OP, FUND_RES, 5_000_000, "انتقال نمونه به اندوخته")]
TR_CAP = 10
INC = 1000
COMMON_WATER_BASIS = B_EQUAL

# --------------------------------------------------------------------------
# Independent reference implementation (exact rationals)
# --------------------------------------------------------------------------

def largest_remainder(amount, weights, inc):
    """weights: list of (unit, sort, weight). Mirrors domain/money/Allocator.kt."""
    total = sum(w for _, _, w in weights)
    if total == 0:
        return None
    floors, rems = {}, []
    for u, s, w in weights:
        num = amount * w
        den = total * inc
        q = num // den
        floors[u] = q * inc
        rems.append((num - den * q, s, u))
    left = amount - sum(floors.values())
    k, residue = divmod(left, inc)
    ranked = sorted(rems, key=lambda r: (-r[0], r[1]))
    for i, (_, _, u) in enumerate(ranked):
        if i < k:
            floors[u] += inc
        if i == 0:
            floors[u] += residue
    assert sum(floors.values()) == amount
    return floors


def reference():
    sort = {u[0]: i + 1 for i, u in enumerate(UNITS)}
    occ = {u[0]: u[4] for u in UNITS}
    area = {u[0]: u[3] for u in UNITS}
    litres = {u: 0 for u in UNIT_IDS}
    main = 0
    for m in METERS:
        prev, cur = round(m[3] * 1000), round(m[4] * 1000)
        cons = cur - prev if cur >= prev else (m[6] * 1000 - prev + cur)
        if m[1] == ROLE_MAIN:
            main += cons
        else:
            litres[m[2]] += cons
    park_all = {u: sum(1 for p in PARKING if p[1] == u) for u in UNIT_IDS}
    park_use = {u: sum(1 for p in PARKING if p[1] == u and p[3] == YES) for u in UNIT_IDS}
    lines = []
    for e in EXPENSES:
        eid, amount, basis = e[0], e[5], e[6]
        if basis == B_NONE:
            continue
        if basis == B_CONS_COMMON:
            sum_units = sum(litres.values())
            metered = int((Fraction(amount * sum_units, main * INC) + Fraction(1, 2)) // 1) * INC
            lines.append((eid + "-a", eid, metered, B_CONS, e[7], e[3]))
            lines.append((eid + "-b", eid, amount - metered, COMMON_WATER_BASIS, e[7], e[3]))
        else:
            lines.append((eid, eid, amount, basis, e[7], e[3]))
    for c in CHARGE_ITEMS:
        lines.append((c[0], c[0], c[4], c[5], c[3], c[2]))
    result = {}
    for lid, src, amount, basis, fund, cat in lines:
        qty = {
            B_EQUAL: {u: 1 for u in UNIT_IDS}, B_OCC: occ, B_AREA: area, B_CONS: litres,
            B_WEIGHT: {u: WEIGHTS.get(src, {}).get(u, 0) for u in UNIT_IDS},
            B_EXPLICIT: {u: WEIGHTS.get(src, {}).get(u, 0) for u in UNIT_IDS},
            B_PARK_ALL: park_all, B_PARK_USE: park_use,
        }[basis]
        shares = largest_remainder(amount, [(u, sort[u], qty[u]) for u in UNIT_IDS], INC)
        result[lid] = {"source": src, "amount": amount, "basis": basis, "fund": fund,
                       "category": cat, "shares": shares}
    totals = {u: sum(l["shares"][u] for l in result.values()) for u in UNIT_IDS}
    pays = {u: sum(p[2] for p in PAYMENTS if p[1] == u) for u in UNIT_IDS}
    credit = {u: sum(a[3] for a in ADJUSTMENTS if a[1] == u and a[2] == ADJ_CREDIT) for u in UNIT_IDS}
    debit = {u: sum(a[3] for a in ADJUSTMENTS if a[1] == u and a[2] == ADJ_DEBIT) for u in UNIT_IDS}
    refund = {u: sum(a[3] for a in ADJUSTMENTS if a[1] == u and a[2] == ADJ_REFUND) for u in UNIT_IDS}
    closing = {u: OPENING["units"].get(u, 0) + totals[u] + debit[u] - credit[u] - pays[u] + refund[u]
               for u in UNIT_IDS}
    cash_close = (OPENING["cash"] + sum(pays.values()) - sum(p[2] for p in SUPPLIER_PAYMENTS)
                  - sum(refund.values()))
    exp_total = sum(e[5] for e in EXPENSES)
    payable_close = OPENING["payable"] + exp_total - sum(p[2] for p in SUPPLIER_PAYMENTS)
    def fund_total(f):
        charges = sum(l["amount"] for l in result.values() if l["fund"] == f)
        adj = sum((a[3] if a[2] == ADJ_DEBIT else -a[3]) for a in ADJUSTMENTS
                  if a[4] == f and a[2] != ADJ_REFUND)
        exps = sum(e[5] for e in EXPENSES if e[7] == f)
        tr = sum(t[3] for t in TRANSFERS if t[2] == f) - sum(t[3] for t in TRANSFERS if t[1] == f)
        return charges + adj - exps + tr
    op_open = OPENING["cash"] + sum(OPENING["units"].values()) - OPENING["payable"] - OPENING["reserve"]
    funds = {FUND_OP: op_open + fund_total(FUND_OP), FUND_RES: OPENING["reserve"] + fund_total(FUND_RES)}
    assert cash_close + sum(closing.values()) - payable_close == funds[FUND_OP] + funds[FUND_RES]
    return {
        "roundingIncrement": INC,
        "waterLitres": {"main": main, "units": litres, "unallocated": main - sum(litres.values())},
        "lines": result, "unitChargeTotals": totals, "unitClosingBalances": closing,
        "cashClosing": cash_close, "payableClosing": payable_close, "funds": funds,
        "operatingOpening": op_open,
    }


# --------------------------------------------------------------------------
# Workbook helpers
# --------------------------------------------------------------------------

def q(sheet):
    return "'" + sheet + "'"


def sheet(wb, title, heading, note=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws["A1"] = heading
    ws["A1"].font = F_TITLE
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE
    return ws


def header(ws, row, names, widths=None):
    for i, n in enumerate(names, start=1):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def inp(ws, ref, value=None, fmt=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font, c.fill, c.border = F_INPUT, FILL_INPUT, BORDER
    if fmt:
        c.number_format = fmt
    return c


def calc(ws, ref, formula, fmt=None, link=False):
    c = ws[ref]
    c.value = formula
    c.font, c.fill, c.border = (F_LINK if link else F_CALC), FILL_CALC, BORDER
    if fmt:
        c.number_format = fmt
    return c


def dv_list(ws, rng, source, allow_blank=True):
    dv = DataValidation(type="list", formula1=source, allow_blank=allow_blank,
                        showErrorMessage=True, errorTitle="مقدار نامعتبر",
                        error="لطفاً از فهرست انتخاب کنید.")
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_nonneg(ws, rng, decimal=False):
    dv = DataValidation(type="decimal" if decimal else "whole", operator="greaterThanOrEqual",
                        formula1="0", allow_blank=True, showErrorMessage=True,
                        errorTitle="مقدار نامعتبر", error="مقدار باید عدد نامنفی باشد.")
    ws.add_data_validation(dv)
    dv.add(rng)


# --------------------------------------------------------------------------
# Build
# --------------------------------------------------------------------------

def build(out):
    ref = reference()
    wb = Workbook()
    wb.remove(wb.active)

    # ---- Guide ----------------------------------------------------------
    g = sheet(wb, "راهنما", "دفتر حسابداری ساختمان «دنگ» — نمونه اولیه (Prototype)")
    rows = [
        "این کارپوشه نمونه اولیه قواعد مالی برنامه اندرویدی «دنگ» است و تمام داده‌های آن ساختگی (نمونه) هستند.",
        "موتور محاسبه اصلی در برنامه (کاتلین) پیاده‌سازی و آزمون شده است؛ این فایل برای اعتبارسنجی قواعد است.",
        "",
        "راهنمای رنگ‌ها:",
        "• سلول زرد با متن آبی = ورودی قابل ویرایش",
        "• سلول سفید با متن سیاه = محاسبه با فرمول (ویرایش نکنید)",
        "• متن سبز = پیوند به برگه دیگر",
        "• برگه «کنترل‌ها» باید همه موارد را «درست» نشان دهد.",
        "",
        "قواعد کلیدی:",
        "• مبالغ به کوچک‌ترین واحد پول (ریال) و به صورت عدد صحیح ثبت می‌شوند.",
        "• سهم هر واحد = مبلغ × وزن واحد ÷ مجموع وزن‌ها؛ سپس به مضرب «واحد گرد کردن» کاهش می‌یابد.",
        "• باقیمانده گرد کردن به روش «بزرگ‌ترین باقیمانده» توزیع می‌شود (تساوی: ترتیب واحد)؛ هیچ باقیمانده‌ای حذف نمی‌شود.",
        "• مصرف آب = قرائت فعلی − قرائت قبلی. مقدار منفی فقط با علامت «گردش کنتور» پذیرفته می‌شود.",
        "• تعویض کنتور: کنتور قدیم با قرائت نهایی و کنتور جدید با قرائت اولیه، هر دو به واحد تعلق دارند.",
        "• اختلاف کنتور اصلی و مجموع کنتورهای واحدها «مصرف تخصیص‌نیافته» است و فقط با سیاست صریح «مصرف + سهم مشترک» تسهیم می‌شود.",
        "• دریافت از ساکن درآمد نیست؛ بدهی واحد را کاهش می‌دهد. انتقال به اندوخته هزینه نیست.",
        "• علامت مانده واحد: مثبت = بدهکار (واحد بدهی دارد)، منفی = بستانکار (اعتبار).",
        "• هزینه «بدون تسهیم» از محل صندوق پرداخت و به واحدها شارژ نمی‌شود.",
        "",
        "محدودیت‌های این نمونه: یک دوره، یک صندوق نقد، تعداد ساکنان در کل دوره ثابت فرض شده است.",
        "برنامه اصلی تاریخچه سکونت، چند دوره، تعرفه پلکانی آب و اصلاحیه پس از قطعی شدن را پشتیبانی می‌کند.",
    ]
    for i, t in enumerate(rows, start=3):
        g.cell(row=i, column=1, value=t).font = Font(name=FONT, size=10)
    inp(g, "A26", "نمونه سلول ورودی")
    calc(g, "B26", "=1+1")
    g.column_dimensions["A"].width = 110
    g.column_dimensions["B"].width = 18

    # ---- Lists ----------------------------------------------------------
    L = sheet(wb, "فهرست‌ها", "فهرست‌های مرجع (برای اعتبارسنجی ورودی‌ها)")
    lists = {"A": ("مبنای تسهیم", BASES), "C": ("بله/خیر", [YES, NO]), "E": ("صندوق", FUNDS),
             "G": ("نوع رابطه پارکینگ", RELATIONS), "I": ("نوع تعدیل", ADJ_TYPES),
             "K": ("نقش کنتور", [ROLE_MAIN, ROLE_UNIT]), "M": ("روش پرداخت", METHODS)}
    for col, (title, vals) in lists.items():
        L[f"{col}3"] = title
        L[f"{col}3"].font = Font(name=FONT, bold=True)
        L.column_dimensions[col].width = 24
        for i, v in enumerate(vals, start=4):
            L[f"{col}{i}"] = v
            L[f"{col}{i}"].font = F_CALC
    LST = q("فهرست‌ها")
    R_BASIS = f"{LST}!$A$4:$A${3 + len(BASES)}"
    R_YN = f"{LST}!$C$4:$C$5"
    R_FUND = f"{LST}!$E$4:$E$5"
    R_REL = f"{LST}!$G$4:$G$6"
    R_ADJ = f"{LST}!$I$4:$I$6"
    R_ROLE = f"{LST}!$K$4:$K$5"
    R_METHOD = f"{LST}!$M$4:$M${3 + len(METHODS)}"

    # ---- Settings -------------------------------------------------------
    S = sheet(wb, "تنظیمات", "تنظیمات ساختمان و دوره")
    settings = [
        ("نام ساختمان", "ساختمان نمونه ۹ واحدی", None),
        ("نشانی", "نشانی نمونه — ساختگی", None),
        ("واحد پول", "ریال", None),
        ("واحد گرد کردن سهم‌ها (ریال)", INC, INT),
        ("نام دوره", "مهر ۱۴۰۵ (نمونه)", None),
        ("شروع دوره (شمسی)", "1405/07/01", None),
        ("پایان دوره (شمسی)", "1405/07/30", None),
        ("تقویم", "هجری شمسی", None),
        ("مبنای تسهیم سهم مشترک آب", COMMON_WATER_BASIS, None),
    ]
    header(S, 3, ["عنوان", "مقدار"], [36, 34])
    for i, (k, v, fmt) in enumerate(settings, start=4):
        S.cell(row=i, column=1, value=k).font = F_CALC
        inp(S, f"B{i}", v, fmt)
    S["B7"].comment = Comment("مثلاً ۱۰۰۰ یعنی سهم‌ها به مضرب ۱۰۰۰ ریال گرد می‌شوند. باقیمانده به روش بزرگ‌ترین باقیمانده توزیع می‌شود.", "Dang")
    dv_list(S, "B12", R_BASIS)
    dv_nonneg(S, "B7")
    INC_REF = f"{q('تنظیمات')}!$B$7"
    COMMON_BASIS_REF = f"{q('تنظیمات')}!$B$12"

    # ---- Units ----------------------------------------------------------
    U = sheet(wb, "واحدها", "واحدها و ساکنان", "تعداد ساکنان «نمونه» است و باید توسط مدیر ساختمان اصلاح شود.")
    header(U, 4, ["شناسه واحد", "طبقه", "شماره واحد", "متراژ (m²)", "تعداد ساکنان (نمونه)", "فعال",
                  "یادداشت", "ترتیب", "مصرف آب دوره (m³)", "مصرف (لیتر)", "پارکینگ تخصیصی", "پارکینگ در استفاده"],
           [12, 8, 10, 12, 14, 8, 24, 8, 16, 14, 14, 14])
    u0 = 5
    uN = u0 + len(UNITS) - 1
    for i, (uid, floor, num, area, occ) in enumerate(UNITS):
        r = u0 + i
        inp(U, f"A{r}", uid)
        inp(U, f"B{r}", floor, INT)
        inp(U, f"C{r}", num, INT)
        inp(U, f"D{r}", area, '#,##0.00')
        inp(U, f"E{r}", occ, INT)
        inp(U, f"F{r}", YES)
        inp(U, f"G{r}", "")
        calc(U, f"H{r}", f"=ROW()-{u0 - 1}", INT)
        W = q("کنتورها")
        calc(U, f"I{r}", f"=SUMIFS({W}!$I$5:$I$30,{W}!$C$5:$C$30,A{r},{W}!$B$5:$B$30,\"{ROLE_UNIT}\")", M3, link=True)
        calc(U, f"J{r}", f"=ROUND(I{r}*1000,0)", INT)
        P = q("پارکینگ")
        calc(U, f"K{r}", f"=COUNTIFS({P}!$B$5:$B$30,A{r})", INT, link=True)
        calc(U, f"L{r}", f"=COUNTIFS({P}!$B$5:$B$30,A{r},{P}!$D$5:$D$30,\"{YES}\")", INT, link=True)
    t = uN + 1
    U[f"A{t}"] = "جمع"
    U[f"A{t}"].font = Font(name=FONT, bold=True)
    for col, fmt in (("D", '#,##0.00'), ("E", INT), ("I", M3), ("J", INT), ("K", INT), ("L", INT)):
        calc(U, f"{col}{t}", f"=SUM({col}{u0}:{col}{uN})", fmt)
    dv_list(U, f"F{u0}:F{uN}", R_YN, False)
    dv_nonneg(U, f"D{u0}:E{uN}", decimal=True)
    UN = q("واحدها")
    R_UNITS = f"{UN}!$A${u0}:$A${uN}"

    # ---- Parking --------------------------------------------------------
    P = sheet(wb, "پارکینگ", "فهرست و تخصیص پارکینگ",
              "تخصیص‌ها ساختگی‌اند. مالکیت/تخصیص از «استفاده واقعی» جداست؛ مسئولیت هزینه با سیاست هر هزینه تعیین می‌شود.")
    header(P, 4, ["شناسه جای پارک", "واحد", "نوع رابطه", "در استفاده", "خودرو / یادداشت"], [16, 10, 14, 12, 28])
    for i, (sid, unit, rel, use, veh) in enumerate(PARKING):
        r = 5 + i
        inp(P, f"A{r}", sid)
        inp(P, f"B{r}", unit or None)
        inp(P, f"C{r}", rel or None)
        inp(P, f"D{r}", use)
        inp(P, f"E{r}", veh)
    for r in range(5 + len(PARKING), 31):
        for col in "ABCDE":
            inp(P, f"{col}{r}")
    dv_list(P, "B5:B30", R_UNITS)
    dv_list(P, "C5:C30", R_REL)
    dv_list(P, "D5:D30", R_YN)

    # ---- Meters ---------------------------------------------------------
    M = sheet(wb, "کنتورها", "قرائت کنتورهای آب (m³)", "قرائت‌ها ساختگی‌اند. مصرف منفی بدون علامت گردش کنتور «خطا» است و در کنترل‌ها گزارش می‌شود.")
    header(M, 4, ["شناسه کنتور", "نقش", "واحد", "قرائت قبلی", "قرائت فعلی", "گردش کنتور؟", "حد گردش (m³)",
                  "یادداشت / تعویض", "مصرف (m³)", "وضعیت"], [14, 8, 8, 14, 14, 12, 12, 34, 12, 18])
    for r in range(5, 31):
        idx = r - 5
        m = METERS[idx] if idx < len(METERS) else None
        inp(M, f"A{r}", m[0] if m else None)
        inp(M, f"B{r}", m[1] if m else None)
        inp(M, f"C{r}", (m[2] or None) if m else None)
        inp(M, f"D{r}", m[3] if m else None, M3)
        inp(M, f"E{r}", m[4] if m else None, M3)
        inp(M, f"F{r}", (m[5] if m else None))
        inp(M, f"G{r}", (m[6] if m else None), INT)
        inp(M, f"H{r}", (m[7] if m else None))
        calc(M, f"I{r}", f'=IF(OR(D{r}="",E{r}=""),0,IF(E{r}>=D{r},E{r}-D{r},IF(AND(F{r}="{YES}",G{r}>0),G{r}-D{r}+E{r},0)))', M3)
        calc(M, f"J{r}", f'=IF(A{r}="","",IF(OR(D{r}="",E{r}=""),"قرائت ناقص",IF(E{r}>=D{r},"عادی",IF(AND(F{r}="{YES}",G{r}>0),"گردش کنتور","خطا: مصرف منفی"))))')
    dv_list(M, "B5:B30", R_ROLE)
    dv_list(M, "C5:C30", R_UNITS)
    dv_list(M, "F5:F30", R_YN)
    dv_nonneg(M, "D5:E30", decimal=True)
    M["L4"] = "تطبیق کنتور اصلی"
    M["L4"].font = Font(name=FONT, bold=True)
    recon = [
        ("مصرف کنتور اصلی (m³)", f'=SUMIFS(I5:I30,B5:B30,"{ROLE_MAIN}")'),
        ("مجموع مصرف کنتورهای واحد (m³)", f'=SUMIFS(I5:I30,B5:B30,"{ROLE_UNIT}")'),
        ("مصرف تخصیص‌نیافته (m³)", "=M5-M6"),
        ("درصد تخصیص‌نیافته", "=IF(M5=0,0,M7/M5)"),
        ("تعداد خطای قرائت", '=COUNTIF(J5:J30,"خطا*")'),
        ("تعداد قرائت ناقص", '=COUNTIF(J5:J30,"قرائت ناقص")'),
    ]
    for i, (k, f) in enumerate(recon, start=5):
        M[f"L{i}"] = k
        M[f"L{i}"].font = F_CALC
        calc(M, f"M{i}", f, "0.0%" if "درصد" in k else (INT if "تعداد" in k else M3))
    M.column_dimensions["L"].width = 30
    M.column_dimensions["M"].width = 14
    MT = q("کنتورها")
    MAIN_REF, UNITSUM_REF = f"{MT}!$M$5", f"{MT}!$M$6"

    # ---- Periods --------------------------------------------------------
    PR = sheet(wb, "دوره‌ها", "دوره‌های صورتحساب")
    header(PR, 4, ["شناسه دوره", "نام", "شروع", "پایان", "مهلت پرداخت", "وضعیت"], [12, 22, 12, 12, 12, 14])
    for col, v in zip("ABCDEF", ["1405-07", "مهر ۱۴۰۵ (نمونه)", "1405/07/01", "1405/07/30", "1405/08/10", "پیش‌نویس"]):
        inp(PR, f"{col}5", v)
    PR["A7"] = "وضعیت «قطعی» یعنی ورودی‌ها، قواعد و نتایج دوره ثابت می‌شوند و اصلاح فقط با سند تعدیل مرتبط انجام می‌شود."
    PR["A7"].font = F_NOTE
    dv_list(PR, "F5:F20", '"پیش‌نویس,قطعی"')

    # ---- Categories -----------------------------------------------------
    C = sheet(wb, "دسته‌ها", "دسته‌های هزینه و سیاست پیش‌فرض تسهیم", "نمونه‌اند، نه قاعده حقوقی. هر ساختمان سیاست خود را انتخاب می‌کند.")
    header(C, 4, ["دسته", "مبنای پیش‌فرض", "صندوق پیش‌فرض", "مربوط به آب؟", "فعال"], [26, 22, 14, 12, 8])
    for i, (name, basis, fund, water) in enumerate(CATEGORIES):
        r = 5 + i
        inp(C, f"A{r}", name)
        inp(C, f"B{r}", basis)
        inp(C, f"C{r}", fund)
        inp(C, f"D{r}", water)
        inp(C, f"E{r}", YES)
    c_end = 5 + len(CATEGORIES) + 5
    for r in range(5 + len(CATEGORIES), c_end + 1):
        for col in "ABCDE":
            inp(C, f"{col}{r}")
    dv_list(C, f"B5:B{c_end}", R_BASIS)
    dv_list(C, f"C5:C{c_end}", R_FUND)
    dv_list(C, f"D5:E{c_end}", R_YN)
    R_CATS = f"{q('دسته‌ها')}!$A$5:$A${c_end}"

    # ---- Expenses -------------------------------------------------------
    E = sheet(wb, "هزینه‌ها", "ثبت هزینه‌ها (تعهد) — پرداخت به تامین‌کننده جداگانه ثبت می‌شود")
    header(E, 4, ["شناسه", "تاریخ", "شرح", "دسته", "تامین‌کننده", "مبلغ (ریال)", "مبنای تسهیم", "صندوق", "دوره",
                  "پرداخت‌شده", "مانده بدهی", "وضعیت پرداخت", "مرجع سند", "یادداشت"],
           [8, 12, 30, 22, 20, 14, 22, 10, 10, 14, 14, 14, 12, 20])
    e0, eN = 5, 5 + EXP_CAP - 1
    SPS = q("پرداخت‌ها")
    for i in range(EXP_CAP):
        r = e0 + i
        e = EXPENSES[i] if i < len(EXPENSES) else None
        inp(E, f"A{r}", f"E{i + 1}")
        for col, val, fmt in (("B", e and e[1], None), ("C", e and e[2], None), ("D", e and e[3], None),
                              ("E", e and (e[4] or None), None), ("F", e and e[5], MONEY),
                              ("G", e and e[6], None), ("H", e and e[7], None),
                              ("I", "1405-07" if e else None, None), ("M", None, None), ("N", None, None)):
            inp(E, f"{col}{r}", val or None, fmt)
        calc(E, f"J{r}", f"=SUMIFS({SPS}!$C$5:$C${4 + SP_CAP},{SPS}!$B$5:$B${4 + SP_CAP},A{r})", MONEY, link=True)
        calc(E, f"K{r}", f"=N(F{r})-J{r}", MONEY)
        calc(E, f"L{r}", f'=IF(N(F{r})=0,"",IF(J{r}=0,"پرداخت‌نشده",IF(J{r}<F{r},"پرداخت جزئی",IF(J{r}=F{r},"پرداخت‌شده","پرداخت مازاد!"))))')
    t = eN + 1
    E[f"C{t}"] = "جمع"
    for col in "FJK":
        calc(E, f"{col}{t}", f"=SUM({col}{e0}:{col}{eN})", MONEY)
    dv_list(E, f"D{e0}:D{eN}", R_CATS)
    dv_list(E, f"G{e0}:G{eN}", R_BASIS)
    dv_list(E, f"H{e0}:H{eN}", R_FUND)
    dv_nonneg(E, f"F{e0}:F{eN}")
    EX = q("هزینه‌ها")

    # ---- Charge schedule ------------------------------------------------
    CH = sheet(wb, "اقلام شارژ", "اقلام شارژ دوره‌ای (مثل حق اندوخته)", "این اقلام هزینه نیستند؛ مستقیماً به واحدها شارژ می‌شوند.")
    header(CH, 4, ["شناسه", "عنوان", "دسته", "صندوق", "مبلغ کل (ریال)", "مبنای تسهیم", "فعال"], [8, 24, 16, 10, 16, 22, 8])
    c0, cN = 5, 5 + CHG_CAP - 1
    for i in range(CHG_CAP):
        r = c0 + i
        c = CHARGE_ITEMS[i] if i < len(CHARGE_ITEMS) else None
        inp(CH, f"A{r}", f"C{i + 1}")
        inp(CH, f"B{r}", c and c[1])
        inp(CH, f"C{r}", c and c[2])
        inp(CH, f"D{r}", c and c[3])
        inp(CH, f"E{r}", c and c[4], MONEY)
        inp(CH, f"F{r}", c and c[5])
        inp(CH, f"G{r}", YES if c else None)
    dv_list(CH, f"C{c0}:C{cN}", R_CATS)
    dv_list(CH, f"D{c0}:D{cN}", R_FUND)
    dv_list(CH, f"F{c0}:F{cN}", R_BASIS)
    dv_list(CH, f"G{c0}:G{cN}", R_YN)
    CHS = q("اقلام شارژ")

    # ---- Inclusion & weights matrices ----------------------------------
    sources = [f"E{i + 1}" for i in range(EXP_CAP)] + [f"C{i + 1}" for i in range(CHG_CAP)]

    def matrix(title, heading, note, default_fn, fmt):
        ws = sheet(wb, title, heading, note)
        header(ws, 4, ["منبع"] + UNIT_IDS, [10] + [12] * len(UNIT_IDS))
        for i, src in enumerate(sources):
            r = 5 + i
            ws[f"A{r}"] = src
            ws[f"A{r}"].font = Font(name=FONT, bold=True)
            for j, uid in enumerate(UNIT_IDS):
                inp(ws, f"{get_column_letter(2 + j)}{r}", default_fn(src, uid), fmt)
            last = get_column_letter(1 + len(UNIT_IDS))
            calc(ws, f"{get_column_letter(2 + len(UNIT_IDS))}{r}", f"=SUM(B{r}:{last}{r})", fmt)
        ws.cell(row=4, column=2 + len(UNIT_IDS), value="جمع ردیف").font = F_HEAD
        ws.cell(row=4, column=2 + len(UNIT_IDS)).fill = FILL_HEAD
        dv_nonneg(ws, f"B5:{get_column_letter(1 + len(UNIT_IDS))}{4 + len(sources)}")
        return ws

    matrix("شمول", "شمول واحدها در هر هزینه/قلم (۱ = مشمول، ۰ = مستثنا)",
           "برای تسهیم میان زیرمجموعه‌ای از واحدها، واحدهای غیرمشمول را ۰ کنید.", lambda s, u: 1, INT)
    matrix("وزن‌ها", "وزن یا مبلغ صریح هر واحد (فقط برای مبنای «وزنی» و «مبلغ صریح»)",
           "وزن‌های آسانسور نمونه‌اند (طبقه اول ۰). در «مبلغ صریح» جمع ردیف باید برابر مبلغ هزینه باشد.",
           lambda s, u: WEIGHTS.get(s, {}).get(u, 0), INT)
    INCL, WGT = q("شمول"), q("وزن‌ها")

    # ---- Allocation lines -----------------------------------------------
    AL = sheet(wb, "سطرهای تسهیم", "سطرهای تسهیم (مشتق از هزینه‌ها و اقلام شارژ)",
               "هر هزینه یک سطر اصلی و یک سطر «سهم مشترک آب» دارد که فقط برای مبنای «مصرف + سهم مشترک» مقدار می‌گیرد.")
    header(AL, 4, ["سطر", "منبع", "شرح", "مبنا", "مبلغ قابل تسهیم", "صندوق", "دسته", "ردیف ماتریس", "مجموع وزن"],
           [8, 8, 34, 22, 16, 10, 22, 10, 14])
    lines = []
    r = 5
    for i in range(EXP_CAP):
        er = e0 + i
        src_row = 5 + i
        lines.append((r, f"E{i + 1}"))
        calc(AL, f"A{r}", f"L{len(lines)}")
        calc(AL, f"B{r}", f"={EX}!A{er}", link=True)
        calc(AL, f"C{r}", f'=IF({EX}!C{er}="","",{EX}!C{er})', link=True)
        calc(AL, f"D{r}", f'=IF({EX}!G{er}="{B_CONS_COMMON}","{B_CONS}",IF({EX}!G{er}="","{B_NONE}",{EX}!G{er}))')
        calc(AL, f"E{r}", (f'=IF(OR({EX}!G{er}="{B_NONE}",{EX}!G{er}=""),0,IF({EX}!G{er}="{B_CONS_COMMON}",'
                           f'IF({MAIN_REF}=0,0,ROUND(N({EX}!F{er})*{UNITSUM_REF}*1000/({MAIN_REF}*1000*{INC_REF}),0)*{INC_REF}),N({EX}!F{er})))'), MONEY)
        calc(AL, f"F{r}", f"={EX}!H{er}", link=True)
        calc(AL, f"G{r}", f"={EX}!D{er}", link=True)
        calc(AL, f"H{r}", src_row, INT)
        r += 1
        lines.append((r, f"E{i + 1}"))
        calc(AL, f"A{r}", f"L{len(lines)}")
        calc(AL, f"B{r}", f"={EX}!A{er}", link=True)
        calc(AL, f"C{r}", f'=IF({EX}!G{er}="{B_CONS_COMMON}","سهم مشترک آب: "&{EX}!C{er},"")')
        calc(AL, f"D{r}", f"={COMMON_BASIS_REF}", link=True)
        calc(AL, f"E{r}", f'=IF({EX}!G{er}="{B_CONS_COMMON}",N({EX}!F{er})-E{r - 1},0)', MONEY)
        calc(AL, f"F{r}", f"={EX}!H{er}", link=True)
        calc(AL, f"G{r}", f"={EX}!D{er}", link=True)
        calc(AL, f"H{r}", src_row, INT)
        r += 1
    for i in range(CHG_CAP):
        cr = c0 + i
        lines.append((r, f"C{i + 1}"))
        calc(AL, f"A{r}", f"L{len(lines)}")
        calc(AL, f"B{r}", f"={CHS}!A{cr}", link=True)
        calc(AL, f"C{r}", f'=IF({CHS}!B{cr}="","",{CHS}!B{cr})', link=True)
        calc(AL, f"D{r}", f'=IF({CHS}!F{cr}="","{B_NONE}",{CHS}!F{cr})', link=True)
        calc(AL, f"E{r}", f'=IF({CHS}!G{cr}="{YES}",N({CHS}!E{cr}),0)', MONEY)
        calc(AL, f"F{r}", f"={CHS}!D{cr}", link=True)
        calc(AL, f"G{r}", f"={CHS}!C{cr}", link=True)
        calc(AL, f"H{r}", EXP_CAP + 5 + i, INT)
        r += 1
    alN = r - 1
    ALS = q("سطرهای تسهیم")

    # ---- Allocation long table -----------------------------------------
    A = sheet(wb, "محاسبه تسهیم", "محاسبه سهم هر واحد — روش بزرگ‌ترین باقیمانده",
              "باقیمانده (عدد صحیح) = مبلغ×وزن − مجموع‌وزن×واحدگرد×⌊مبلغ×وزن÷(مجموع‌وزن×واحدگرد)⌋ ؛ رتبه بر اساس باقیمانده نزولی و سپس ترتیب واحد.")
    header(A, 4, ["سطر", "منبع", "واحد", "ترتیب", "مبنا", "مبلغ سطر", "مشمول", "کمیت مبنا", "وزن", "مجموع وزن",
                  "سهم پایه", "باقیمانده", "رتبه", "باقیمانده کل", "تعداد افزایش", "خرده نهایی", "سهم نهایی", "صندوق", "دسته", "آب؟"],
           [7, 7, 7, 6, 20, 13, 7, 11, 11, 12, 13, 14, 6, 12, 8, 9, 13, 9, 20, 7])
    row = 5
    a0 = row
    n_rows = len(lines) * len(UNIT_IDS)
    aN = a0 + n_rows - 1
    rngA = f"$A${a0}:$A${aN}"
    rngD = f"$D${a0}:$D${aN}"
    rngK = f"$K${a0}:$K${aN}"
    rngL = f"$L${a0}:$L${aN}"
    rngI = f"$I${a0}:$I${aN}"
    for li, (lrow, src) in enumerate(lines):
        for ui, uid in enumerate(UNIT_IDS):
            r = row
            ur = u0 + ui
            calc(A, f"A{r}", f"={ALS}!A{lrow}", link=True)
            calc(A, f"B{r}", f"={ALS}!B{lrow}", link=True)
            calc(A, f"C{r}", f"={UN}!A{ur}", link=True)
            calc(A, f"D{r}", f"={UN}!H{ur}", INT, link=True)
            calc(A, f"E{r}", f"={ALS}!D{lrow}", link=True)
            calc(A, f"F{r}", f"={ALS}!E{lrow}", MONEY, link=True)
            mrow = 5 + (int(src[1:]) - 1 if src[0] == "E" else EXP_CAP + int(src[1:]) - 1)
            mcol = get_column_letter(2 + ui)
            calc(A, f"G{r}", f'=IF({UN}!F{ur}="{YES}",{INCL}!{mcol}{mrow},0)', INT, link=True)
            calc(A, f"H{r}", (f'=CHOOSE(MATCH(E{r},{R_BASIS},0),1,{UN}!E{ur},{UN}!D{ur}*100,{UN}!J{ur},{UN}!J{ur},'
                              f'{WGT}!{mcol}{mrow},{WGT}!{mcol}{mrow},{UN}!K{ur},{UN}!L{ur},0)'), INT)
            calc(A, f"I{r}", f"=G{r}*H{r}", INT)
            calc(A, f"J{r}", f"=SUMIFS({rngI},{rngA},A{r})", INT)
            calc(A, f"K{r}", f"=IF(J{r}=0,0,{INC_REF}*INT(F{r}*I{r}/(J{r}*{INC_REF})))", MONEY)
            calc(A, f"L{r}", f"=IF(J{r}=0,0,F{r}*I{r}-J{r}*{INC_REF}*INT(F{r}*I{r}/(J{r}*{INC_REF})))", INT)
            calc(A, f"M{r}", f'=1+COUNTIFS({rngA},A{r},{rngL},">"&L{r})+COUNTIFS({rngA},A{r},{rngL},L{r},{rngD},"<"&D{r})', INT)
            calc(A, f"N{r}", f"=IF(J{r}=0,0,F{r}-SUMIFS({rngK},{rngA},A{r}))", MONEY)
            calc(A, f"O{r}", f"=INT(N{r}/{INC_REF})", INT)
            calc(A, f"P{r}", f"=N{r}-O{r}*{INC_REF}", INT)
            calc(A, f"Q{r}", f"=IF(J{r}=0,0,K{r}+IF(M{r}<=O{r},{INC_REF},0)+IF(M{r}=1,P{r},0))", MONEY)
            calc(A, f"R{r}", f"={ALS}!F{lrow}", link=True)
            calc(A, f"S{r}", f"={ALS}!G{lrow}", link=True)
            calc(A, f"T{r}", f'=IFERROR(INDEX({q("دسته‌ها")}!$D$5:$D${c_end},MATCH(S{r},{q("دسته‌ها")}!$A$5:$A${c_end},0)),"{NO}")')
            row += 1
    AT = q("محاسبه تسهیم")
    rngQ = f"{AT}!$Q${a0}:$Q${aN}"
    rngAA = f"{AT}!$A${a0}:$A${aN}"
    rngAC = f"{AT}!$C${a0}:$C${aN}"
    rngAR = f"{AT}!$R${a0}:$R${aN}"
    rngAS = f"{AT}!$S${a0}:$S${aN}"
    rngAI = f"{AT}!$I${a0}:$I${aN}"
    for lrow, _ in lines:
        calc(AL, f"I{lrow}", f"=SUMIFS({rngAI},{rngAA},A{lrow})", INT, link=True)
    rngAT = f"{AT}!$T${a0}:$T${aN}"

    # ---- Charge detail matrix ------------------------------------------
    D = sheet(wb, "ریز شارژ", "ریز شارژ هر واحد به تفکیک سطر تسهیم", "جمع هر سطر باید برابر مبلغ سطر باشد (ستون «اختلاف» = ۰).")
    header(D, 4, ["سطر", "شرح", "مبلغ سطر"] + UNIT_IDS + ["جمع تسهیم", "اختلاف"],
           [7, 34, 14] + [12] * len(UNIT_IDS) + [14, 10])
    d0 = 5
    for li, (lrow, src) in enumerate(lines):
        r = d0 + li
        calc(D, f"A{r}", f"={ALS}!A{lrow}", link=True)
        calc(D, f"B{r}", f"={ALS}!C{lrow}", link=True)
        calc(D, f"C{r}", f"={ALS}!E{lrow}", MONEY, link=True)
        for ui in range(len(UNIT_IDS)):
            col = get_column_letter(4 + ui)
            calc(D, f"{col}{r}", f"=SUMIFS({rngQ},{rngAA},$A{r},{rngAC},{col}$4)", MONEY)
        last = get_column_letter(3 + len(UNIT_IDS))
        tot = get_column_letter(4 + len(UNIT_IDS))
        dif = get_column_letter(5 + len(UNIT_IDS))
        calc(D, f"{tot}{r}", f"=SUM(D{r}:{last}{r})", MONEY)
        calc(D, f"{dif}{r}", f"=C{r}-{tot}{r}", MONEY)
    dN = d0 + len(lines) - 1
    t = dN + 1
    D[f"B{t}"] = "جمع"
    for ci in range(3, 6 + len(UNIT_IDS)):
        col = get_column_letter(ci)
        calc(D, f"{col}{t}", f"=SUM({col}{d0}:{col}{dN})", MONEY)
    DIF_COL = get_column_letter(5 + len(UNIT_IDS))

    # ---- Opening balances -----------------------------------------------
    O = sheet(wb, "افتتاحیه", "مانده‌های افتتاحیه (ساختگی)", "مانده واحد: مثبت = بدهکار، منفی = بستانکار.")
    header(O, 4, ["عنوان", "مبلغ (ریال)"], [34, 18])
    O["A5"], O["A6"], O["A7"] = "موجودی نقد افتتاحیه", "مانده افتتاحیه صندوق اندوخته", "بدهی افتتاحیه به تامین‌کنندگان"
    inp(O, "B5", OPENING["cash"], MONEY)
    inp(O, "B6", OPENING["reserve"], MONEY)
    inp(O, "B7", OPENING["payable"], MONEY)
    O["A8"] = "مانده افتتاحیه صندوق عملیاتی (محاسبه‌شده)"
    calc(O, "B8", f"=B5+SUM(B11:B{10 + len(UNITS)})-B7-B6", MONEY)
    O["B8"].comment = Comment("نقد + مطالبات − بدهی‌ها − اندوخته", "Dang")
    O["A10"] = "مانده افتتاحیه واحدها"
    O["A10"].font = Font(name=FONT, bold=True)
    for i, uid in enumerate(UNIT_IDS):
        r = 11 + i
        calc(O, f"A{r}", f"={UN}!A{u0 + i}", link=True)
        inp(O, f"B{r}", OPENING["units"].get(uid, 0), MONEY)
    OP = q("افتتاحیه")

    # ---- Payments received ---------------------------------------------
    PY = sheet(wb, "دریافت‌ها", "دریافت از واحدها (کاهش بدهی واحد — درآمد نیست)")
    header(PY, 4, ["تاریخ", "واحد", "مبلغ (ریال)", "روش", "مرجع", "یادداشت"], [12, 8, 14, 14, 10, 24])
    for i in range(PAY_CAP):
        r = 5 + i
        p = PAYMENTS[i] if i < len(PAYMENTS) else None
        inp(PY, f"A{r}", p and p[0])
        inp(PY, f"B{r}", p and p[1])
        inp(PY, f"C{r}", p and p[2], MONEY)
        inp(PY, f"D{r}", p and p[3])
        inp(PY, f"E{r}", p and p[4])
        inp(PY, f"F{r}")
    dv_list(PY, f"B5:B{4 + PAY_CAP}", R_UNITS)
    dv_list(PY, f"D5:D{4 + PAY_CAP}", R_METHOD)
    dv_nonneg(PY, f"C5:C{4 + PAY_CAP}")
    PYS = q("دریافت‌ها")

    # ---- Supplier payments ---------------------------------------------
    SP = sheet(wb, "پرداخت‌ها", "پرداخت به تامین‌کنندگان (تسویه بدهی هزینه)", "پرداخت جزئی مجاز است؛ مانده در برگه هزینه‌ها دیده می‌شود.")
    header(SP, 4, ["تاریخ", "شناسه هزینه", "مبلغ (ریال)", "تامین‌کننده", "مرجع"], [12, 12, 14, 22, 10])
    for i in range(SP_CAP):
        r = 5 + i
        s = SUPPLIER_PAYMENTS[i] if i < len(SUPPLIER_PAYMENTS) else None
        inp(SP, f"A{r}", s and s[0])
        inp(SP, f"B{r}", s and s[1])
        inp(SP, f"C{r}", s and s[2], MONEY)
        calc(SP, f"D{r}", f'=IF(B{r}="","",IFERROR(INDEX({EX}!$E${e0}:$E${eN},MATCH(B{r},{EX}!$A${e0}:$A${eN},0))&"","شناسه نامعتبر"))', link=True)
        inp(SP, f"E{r}", s and s[3])
    dv_list(SP, f"B5:B{4 + SP_CAP}", f"{EX}!$A${e0}:$A${eN}")
    dv_nonneg(SP, f"C5:C{4 + SP_CAP}")

    # ---- Adjustments & refunds -----------------------------------------
    AD = sheet(wb, "تعدیل‌ها", "تعدیل‌ها و استرداد وجه", "ذکر دلیل الزامی است. استرداد وجه از صندوق نقد پرداخت و به بدهی واحد افزوده می‌شود.")
    header(AD, 4, ["تاریخ", "واحد", "نوع", "مبلغ (ریال)", "صندوق", "دلیل (الزامی)", "دوره مرتبط"], [12, 8, 22, 14, 10, 40, 10])
    for i in range(ADJ_CAP):
        r = 5 + i
        a = ADJUSTMENTS[i] if i < len(ADJUSTMENTS) else None
        for col, val, fmt in (("A", a and a[0], None), ("B", a and a[1], None), ("C", a and a[2], None),
                              ("D", a and a[3], MONEY), ("E", a and a[4], None), ("F", a and a[5], None),
                              ("G", "1405-07" if a else None, None)):
            inp(AD, f"{col}{r}", val, fmt)
    dv_list(AD, f"B5:B{4 + ADJ_CAP}", R_UNITS)
    dv_list(AD, f"C5:C{4 + ADJ_CAP}", R_ADJ)
    dv_list(AD, f"E5:E{4 + ADJ_CAP}", R_FUND)
    dv_nonneg(AD, f"D5:D{4 + ADJ_CAP}")
    ADS = q("تعدیل‌ها")

    # ---- Transfers --------------------------------------------------------
    TR = sheet(wb, "انتقال‌ها", "انتقال بین صندوق‌ها (هزینه یا درآمد نیست)")
    header(TR, 4, ["تاریخ", "از صندوق", "به صندوق", "مبلغ (ریال)", "یادداشت"], [12, 12, 12, 14, 30])
    for i in range(TR_CAP):
        r = 5 + i
        x = TRANSFERS[i] if i < len(TRANSFERS) else None
        inp(TR, f"A{r}", x and x[0])
        inp(TR, f"B{r}", x and x[1])
        inp(TR, f"C{r}", x and x[2])
        inp(TR, f"D{r}", x and x[3], MONEY)
        inp(TR, f"E{r}", x and x[4])
    dv_list(TR, f"B5:C{4 + TR_CAP}", R_FUND)
    dv_nonneg(TR, f"D5:D{4 + TR_CAP}")
    TRS = q("انتقال‌ها")

    # ---- Cash ledger ------------------------------------------------------
    CL = sheet(wb, "دفتر صندوق", "دفتر صندوق نقد ساختمان (تمام گردش‌های نقدی)",
               "فقط از برگه‌های مبدأ پر می‌شود؛ ویرایش نکنید. انتقال بین صندوق‌ها گردش نقدی ندارد.")
    header(CL, 4, ["ردیف", "تاریخ", "نوع", "طرف حساب", "شرح / مرجع", "ورود (ریال)", "خروج (ریال)", "مانده"],
           [6, 12, 20, 18, 26, 14, 14, 16])
    r = 5
    calc(CL, f"A{r}", 1, INT)
    CL[f"C{r}"] = "مانده افتتاحیه"
    calc(CL, f"F{r}", f"={OP}!B5", MONEY, link=True)
    calc(CL, f"G{r}", 0, MONEY)
    calc(CL, f"H{r}", f"=F{r}-G{r}", MONEY)
    r += 1
    for i in range(PAY_CAP):
        sr = 5 + i
        calc(CL, f"A{r}", f"=A{r - 1}+1", INT)
        calc(CL, f"B{r}", f'=IF({PYS}!C{sr}="","",{PYS}!A{sr})', link=True)
        calc(CL, f"C{r}", f'=IF({PYS}!C{sr}="","","دریافت از واحد")')
        calc(CL, f"D{r}", f'=IF({PYS}!C{sr}="","",{PYS}!B{sr})', link=True)
        calc(CL, f"E{r}", f'=IF({PYS}!C{sr}="","",{PYS}!E{sr}&"")', link=True)
        calc(CL, f"F{r}", f"=N({PYS}!C{sr})", MONEY, link=True)
        calc(CL, f"G{r}", 0, MONEY)
        calc(CL, f"H{r}", f"=H{r - 1}+F{r}-G{r}", MONEY)
        r += 1
    for i in range(SP_CAP):
        sr = 5 + i
        calc(CL, f"A{r}", f"=A{r - 1}+1", INT)
        calc(CL, f"B{r}", f'=IF({SPS}!C{sr}="","",{SPS}!A{sr})', link=True)
        calc(CL, f"C{r}", f'=IF({SPS}!C{sr}="","","پرداخت به تامین‌کننده")')
        calc(CL, f"D{r}", f'=IF({SPS}!C{sr}="","",{SPS}!D{sr})', link=True)
        calc(CL, f"E{r}", f'=IF({SPS}!C{sr}="","",{SPS}!B{sr}&" / "&{SPS}!E{sr})', link=True)
        calc(CL, f"F{r}", 0, MONEY)
        calc(CL, f"G{r}", f"=N({SPS}!C{sr})", MONEY, link=True)
        calc(CL, f"H{r}", f"=H{r - 1}+F{r}-G{r}", MONEY)
        r += 1
    for i in range(ADJ_CAP):
        sr = 5 + i
        cond = f'{ADS}!C{sr}="{ADJ_REFUND}"'
        calc(CL, f"A{r}", f"=A{r - 1}+1", INT)
        calc(CL, f"B{r}", f'=IF({cond},{ADS}!A{sr},"")', link=True)
        calc(CL, f"C{r}", f'=IF({cond},"استرداد وجه به واحد","")')
        calc(CL, f"D{r}", f'=IF({cond},{ADS}!B{sr},"")', link=True)
        calc(CL, f"E{r}", f'=IF({cond},{ADS}!F{sr},"")', link=True)
        calc(CL, f"F{r}", 0, MONEY)
        calc(CL, f"G{r}", f"=IF({cond},N({ADS}!D{sr}),0)", MONEY, link=True)
        calc(CL, f"H{r}", f"=H{r - 1}+F{r}-G{r}", MONEY)
        r += 1
    clN = r - 1
    CL[f"E{r}"] = "جمع / مانده پایانی"
    calc(CL, f"F{r}", f"=SUM(F5:F{clN})", MONEY)
    calc(CL, f"G{r}", f"=SUM(G5:G{clN})", MONEY)
    calc(CL, f"H{r}", f"=F{r}-G{r}", MONEY)
    CASH_CLOSE = f"{q('دفتر صندوق')}!$H${r}"

    # ---- Unit balances ----------------------------------------------------
    B = sheet(wb, "مانده واحدها", "مانده حساب واحدها", "مثبت = بدهکار (واحد بدهی دارد) ؛ منفی = بستانکار (اعتبار واحد).")
    header(B, 4, ["واحد", "مانده افتتاحیه", "شارژ عملیاتی (بدون آب)", "شارژ آب", "شارژ اندوخته", "تعدیل بدهکار",
                  "تعدیل بستانکار", "دریافتی", "استرداد", "مانده پایانی", "وضعیت"],
           [8, 14, 18, 14, 14, 13, 13, 14, 12, 15, 12])
    WATER_CAT = q("دسته‌ها")
    for i, uid in enumerate(UNIT_IDS):
        r = 5 + i
        calc(B, f"A{r}", f"={UN}!A{u0 + i}", link=True)
        calc(B, f"B{r}", f"={OP}!B{11 + i}", MONEY, link=True)
        water_sum = f'SUMIFS({rngQ},{rngAC},A{r},{rngAR},"{FUND_OP}",{rngAT},"{YES}")'
        calc(B, f"C{r}", f'=SUMIFS({rngQ},{rngAC},A{r},{rngAR},"{FUND_OP}")-D{r}', MONEY)
        calc(B, f"D{r}", "=" + water_sum, MONEY)
        calc(B, f"E{r}", f'=SUMIFS({rngQ},{rngAC},A{r},{rngAR},"{FUND_RES}")', MONEY)
        calc(B, f"F{r}", f'=SUMIFS({ADS}!$D$5:$D${4 + ADJ_CAP},{ADS}!$B$5:$B${4 + ADJ_CAP},A{r},{ADS}!$C$5:$C${4 + ADJ_CAP},"{ADJ_DEBIT}")', MONEY, link=True)
        calc(B, f"G{r}", f'=SUMIFS({ADS}!$D$5:$D${4 + ADJ_CAP},{ADS}!$B$5:$B${4 + ADJ_CAP},A{r},{ADS}!$C$5:$C${4 + ADJ_CAP},"{ADJ_CREDIT}")', MONEY, link=True)
        calc(B, f"H{r}", f"=SUMIFS({PYS}!$C$5:$C${4 + PAY_CAP},{PYS}!$B$5:$B${4 + PAY_CAP},A{r})", MONEY, link=True)
        calc(B, f"I{r}", f'=SUMIFS({ADS}!$D$5:$D${4 + ADJ_CAP},{ADS}!$B$5:$B${4 + ADJ_CAP},A{r},{ADS}!$C$5:$C${4 + ADJ_CAP},"{ADJ_REFUND}")', MONEY, link=True)
        calc(B, f"J{r}", f"=B{r}+C{r}+D{r}+E{r}+F{r}-G{r}-H{r}+I{r}", MONEY)
        calc(B, f"K{r}", f'=IF(J{r}>0,"بدهکار",IF(J{r}<0,"بستانکار","تسویه"))')
    bN = 5 + len(UNIT_IDS) - 1
    t = bN + 1
    B[f"A{t}"] = "جمع"
    for col in "BCDEFGHIJ":
        calc(B, f"{col}{t}", f"=SUM({col}5:{col}{bN})", MONEY)
    BAL = q("مانده واحدها")
    RECV_CLOSE = f"{BAL}!$J${t}"

    # ---- Funds --------------------------------------------------------------
    FD = sheet(wb, "صندوق‌ها", "صندوق‌های عملیاتی و اندوخته", "مانده صندوق = افتتاحیه + شارژها ± تعدیل‌ها − هزینه‌ها ± انتقال‌ها")
    header(FD, 4, ["عنوان", FUND_OP, FUND_RES, "جمع"], [34, 16, 16, 16])
    fund_rows = [
        ("مانده افتتاحیه", f"={OP}!B8", f"={OP}!B6"),
        ("شارژ واحدها", f'=SUMIFS({ALS}!$E$5:$E${alN},{ALS}!$F$5:$F${alN},B$4)', None),
        ("تعدیل بدهکار", f'=SUMIFS({ADS}!$D$5:$D${4 + ADJ_CAP},{ADS}!$E$5:$E${4 + ADJ_CAP},B$4,{ADS}!$C$5:$C${4 + ADJ_CAP},"{ADJ_DEBIT}")', None),
        ("تعدیل بستانکار (−)", f'=-SUMIFS({ADS}!$D$5:$D${4 + ADJ_CAP},{ADS}!$E$5:$E${4 + ADJ_CAP},B$4,{ADS}!$C$5:$C${4 + ADJ_CAP},"{ADJ_CREDIT}")', None),
        ("هزینه‌ها (−)", f"=-SUMIFS({EX}!$F${e0}:$F${eN},{EX}!$H${e0}:$H${eN},B$4)", None),
        ("انتقال ورودی", f"=SUMIFS({TRS}!$D$5:$D${4 + TR_CAP},{TRS}!$C$5:$C${4 + TR_CAP},B$4)", None),
        ("انتقال خروجی (−)", f"=-SUMIFS({TRS}!$D$5:$D${4 + TR_CAP},{TRS}!$B$5:$B${4 + TR_CAP},B$4)", None),
    ]
    for i, (k, fop, fres) in enumerate(fund_rows):
        r = 5 + i
        FD[f"A{r}"] = k
        calc(FD, f"B{r}", fop, MONEY)
        calc(FD, f"C{r}", fres if fres else fop.replace("B$4", "C$4"), MONEY)
        calc(FD, f"D{r}", f"=B{r}+C{r}", MONEY)
    fr = 5 + len(fund_rows)
    FD[f"A{fr}"] = "مانده پایانی"
    FD[f"A{fr}"].font = Font(name=FONT, bold=True)
    for col in "BCD":
        calc(FD, f"{col}{fr}", f"=SUM({col}5:{col}{fr - 1})", MONEY)
    FDS = q("صندوق‌ها")

    # ---- Summary -------------------------------------------------------------
    SM = sheet(wb, "خلاصه ماهانه", "خلاصه ماهانه و تطبیق", "بدهی ساکنان (مطالبات) با موجودی نقد ساختمان متفاوت است.")
    SM.column_dimensions["A"].width = 44
    SM.column_dimensions["B"].width = 18
    SM.column_dimensions["D"].width = 30
    SM.column_dimensions["E"].width = 16
    SM.column_dimensions["F"].width = 16
    blocks = [
        ("موجودی نقد", None),
        ("موجودی نقد افتتاحیه", f"={OP}!B5"),
        ("دریافت از واحدها", f"=SUM({PYS}!C5:C{4 + PAY_CAP})"),
        ("پرداخت به تامین‌کنندگان", f"=-SUM({SPS}!C5:C{4 + SP_CAP})"),
        ("استرداد وجه", f'=-SUMIFS({ADS}!$D$5:$D${4 + ADJ_CAP},{ADS}!$C$5:$C${4 + ADJ_CAP},"{ADJ_REFUND}")'),
        ("موجودی نقد پایانی", "=SUM(B5:B8)"),
        ("", None),
        ("وضعیت مالی", None),
        ("مطالبات از واحدها (خالص)", f"={RECV_CLOSE}"),
        ("بدهی به تامین‌کنندگان", f"={OP}!B7+{EX}!F{eN + 1}-SUM({SPS}!C5:C{4 + SP_CAP})"),
        ("مانده صندوق عملیاتی", f"={FDS}!B{fr}"),
        ("مانده صندوق اندوخته", f"={FDS}!C{fr}"),
        ("تطبیق: نقد + مطالبات − بدهی − صندوق‌ها (باید ۰ باشد)", "=B9+B12-B13-B14-B15"),
        ("", None),
        ("آب", None),
        ("مصرف کنتور اصلی (m³)", f"={MAIN_REF}"),
        ("مجموع کنتورهای واحد (m³)", f"={UNITSUM_REF}"),
        ("مصرف تخصیص‌نیافته (m³)", f"={MT}!$M$7"),
    ]
    for i, (k, f) in enumerate(blocks, start=4):
        SM[f"A{i}"] = k
        if f is None:
            SM[f"A{i}"].font = Font(name=FONT, bold=True, size=11)
        else:
            SM[f"A{i}"].font = F_CALC
            calc(SM, f"B{i}", f, M3 if "m³" in k else MONEY, link="!" in f)
    SM["D4"] = "هزینه‌ها به تفکیک دسته"
    SM["D4"].font = Font(name=FONT, bold=True, size=11)
    for col, h in zip("DEF", ["دسته", "هزینه ثبت‌شده", "شارژ تسهیم‌شده"]):
        SM[f"{col}5"] = h
        SM[f"{col}5"].font = F_HEAD
        SM[f"{col}5"].fill = FILL_HEAD
    for i in range(len(CATEGORIES)):
        r = 6 + i
        calc(SM, f"D{r}", f"={q('دسته‌ها')}!A{5 + i}", link=True)
        calc(SM, f"E{r}", f"=SUMIFS({EX}!$F${e0}:$F${eN},{EX}!$D${e0}:$D${eN},D{r})", MONEY, link=True)
        calc(SM, f"F{r}", f"=SUMIFS({rngQ},{rngAS},D{r})", MONEY, link=True)
    sr = 6 + len(CATEGORIES)
    SM[f"D{sr}"] = "جمع"
    calc(SM, f"E{sr}", f"=SUM(E6:E{sr - 1})", MONEY)
    calc(SM, f"F{sr}", f"=SUM(F6:F{sr - 1})", MONEY)

    # ---- Checks & scenario tests --------------------------------------------
    K = sheet(wb, "کنترل‌ها", "کنترل‌های یکپارچگی و سناریوهای آزمون",
              "مقادیر مورد انتظار با پیاده‌سازی مستقل (کسر دقیق) محاسبه و در آزمون کاتلین هم بررسی شده‌اند.")
    header(K, 4, ["کنترل / سناریو", "مقدار محاسبه‌شده", "مقدار مورد انتظار", "نتیجه"], [60, 18, 18, 10])
    checks = [
        ("جمع اختلاف سطرهای تسهیم با مبلغ سطر (باید ۰)", f"=SUMPRODUCT(ABS({q('ریز شارژ')}!{DIF_COL}{d0}:{DIF_COL}{dN}))", 0),
        ("تطبیق ترازنامه (نقد + مطالبات − بدهی − صندوق‌ها)", f"={q('خلاصه ماهانه')}!B16", 0),
        ("مانده نقد دفتر صندوق = خلاصه", f"={CASH_CLOSE}-{q('خلاصه ماهانه')}!B9", 0),
        ("تعداد خطای قرائت کنتور", f"={MT}!$M$9", 0),
        ("تعداد قرائت ناقص", f"={MT}!$M$10", 0),
        ("سهم مشترک آب منفی نیست (۱=درست)", f"=IF({MT}!$M$7>=0,1,0)", 1),
        ("سطر با مبلغ و بدون وزن (تسهیم‌نشده)", f'=COUNTIFS({ALS}!E5:E{alN},">0",{ALS}!I5:I{alN},0)', 0),
        ("مبلغ صریح: اختلاف جمع ردیف با مبلغ هزینه",
         f'=SUMPRODUCT(({EX}!G{e0}:G{eN}="{B_EXPLICIT}")*ABS({EX}!F{e0}:F{eN}-{WGT}!{get_column_letter(2 + len(UNIT_IDS))}5:{get_column_letter(2 + len(UNIT_IDS))}{4 + EXP_CAP}))', 0),
        ("تعدیل بدون دلیل", f'=SUMPRODUCT(({ADS}!D5:D{4 + ADJ_CAP}>0)*({ADS}!F5:F{4 + ADJ_CAP}=""))', 0),
        ("پرداخت مازاد به تامین‌کننده", f'=COUNTIF({EX}!L{e0}:L{eN},"پرداخت مازاد!")', 0),
        ("مصرف تخصیص‌نیافته آب (لیتر)", f"=ROUND({MT}!$M$7*1000,0)", ref["waterLitres"]["unallocated"]),
        ("مانده پایانی نقد", f"={CASH_CLOSE}", ref["cashClosing"]),
        ("مانده صندوق عملیاتی", f"={FDS}!B{fr}", ref["funds"][FUND_OP]),
        ("مانده صندوق اندوخته", f"={FDS}!C{fr}", ref["funds"][FUND_RES]),
        ("بدهی به تامین‌کنندگان", f"={q('خلاصه ماهانه')}!B13", ref["payableClosing"]),
    ]
    for i, uid in enumerate(UNIT_IDS):
        checks.append((f"جمع شارژ واحد {uid}", f"=SUMIFS({rngQ},{rngAC},\"{uid}\")", ref["unitChargeTotals"][uid]))
        checks.append((f"مانده پایانی واحد {uid}", f"={BAL}!J{5 + i}", ref["unitClosingBalances"][uid]))
    sample = [("E1", "3-1"), ("E3", "1-1"), ("E4-a", "1-2"), ("E4-b", "1-1"), ("E5", "3-1"), ("E6", "4-2"), ("E7", "2-1")]
    line_names = {}
    for idx, (lrow, src) in enumerate(lines):
        line_names.setdefault(src, []).append(f"L{idx + 1}")
    for lid, uid in sample:
        src = lid.split("-")[0]
        lname = line_names[src][1 if lid.endswith("-b") else 0]
        checks.append((f"سهم {uid} از {lid} ({ref['lines'][lid]['basis']})",
                       f'=SUMIFS({rngQ},{rngAA},"{lname}",{rngAC},"{uid}")', ref["lines"][lid]["shares"][uid]))
    for i, (k, f, exp) in enumerate(checks):
        r = 5 + i
        K[f"A{r}"] = k
        K[f"A{r}"].font = F_CALC
        calc(K, f"B{r}", f, MONEY)
        c = K.cell(row=r, column=3, value=exp)
        c.font, c.number_format, c.border = F_CALC, MONEY, BORDER
        calc(K, f"D{r}", f'=IF(B{r}=C{r},"درست","نادرست")')
        K[f"D{r}"].fill = FILL_CHECK
    kN = 5 + len(checks) - 1
    K[f"A{kN + 2}"] = "همه کنترل‌ها"
    K[f"A{kN + 2}"].font = Font(name=FONT, bold=True)
    calc(K, f"D{kN + 2}", f'=IF(COUNTIF(D5:D{kN},"درست")={len(checks)},"درست","نادرست")')
    K[f"A{kN + 3}"] = "مقادیر «مورد انتظار» ثابت‌اند و فقط برای همین داده نمونه معتبرند؛ با تغییر ورودی‌ها انتظار می‌رود برخی سناریوها «نادرست» شوند."
    K[f"A{kN + 3}"].font = F_NOTE

    wb.properties.title = "Dang (دنگ) — Building accounting prototype"
    wb.properties.creator = "Dang"
    order = ["راهنما", "تنظیمات", "واحدها", "پارکینگ", "کنتورها", "دوره‌ها", "دسته‌ها", "هزینه‌ها", "اقلام شارژ",
             "شمول", "وزن‌ها", "سطرهای تسهیم", "محاسبه تسهیم", "ریز شارژ", "افتتاحیه", "دریافت‌ها", "پرداخت‌ها",
             "تعدیل‌ها", "انتقال‌ها", "دفتر صندوق", "مانده واحدها", "صندوق‌ها", "خلاصه ماهانه", "کنترل‌ها", "فهرست‌ها"]
    wb._sheets = [wb[n] for n in order]
    wb.save(out)
    serializable = json.loads(json.dumps(ref, ensure_ascii=False))
    Path(out).with_name("expected-results.json").write_text(
        json.dumps(serializable, ensure_ascii=False, indent=2), encoding="utf-8")
    return ref, len(checks)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).with_name("Dang-accounting-prototype.xlsx"))
    ref, n = build(out)
    print(f"wrote {out} with {n} checks")
    print(json.dumps({"unitChargeTotals": ref["unitChargeTotals"], "closing": ref["unitClosingBalances"],
                      "cash": ref["cashClosing"], "funds": ref["funds"], "water": ref["waterLitres"]},
                     ensure_ascii=False, indent=1))
